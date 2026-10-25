"""Offline tests for bounded, authenticated, phone-file research imports."""
import importlib
from io import BytesIO
import os
import sys
import tempfile
import time
import unittest
from unittest.mock import patch

sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))
import file_import


class FileImportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        os.environ["STUDIO_DB_PATH"] = os.path.join(cls.temp.name, "import-test.sqlite3")
        os.environ["STUDIO_ACCESS_TOKEN"] = "file-import-test-token"
        os.environ.pop("STUDIO_ALLOW_API_USAGE", None)
        import app
        cls.app_module = importlib.reload(app)
        cls.client = cls.app_module.app.test_client()
        cls.headers = {"X-Studio-Token": "file-import-test-token"}

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def upload(self, name, data, authenticated=True):
        return self.client.post("/api/files/extract",
            data={"file": (BytesIO(data), name)},
            headers=self.headers if authenticated else {},
            content_type="multipart/form-data")

    def test_requires_owner_authentication_and_does_not_write_document(self):
        self.assertEqual(self.upload("private.txt", b"confidential", False).status_code, 401)
        response = self.upload("notes.txt", b"Meeting notes: test me")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json["name"], "notes.txt")
        self.assertIn("Meeting notes", response.json["text"])
        self.assertFalse(response.json["truncated"])
        self.assertEqual(os.listdir(self.temp.name), ["import-test.sqlite3"])

    def test_rejects_executable_oversize_and_unreadable_documents(self):
        self.assertEqual(self.upload("payload.exe", b"hello").status_code, 422)
        self.assertEqual(self.upload("large.txt", b"a" * (file_import.MAX_FILE_BYTES + 1)).status_code, 422)
        self.assertEqual(self.upload("empty.txt", b"").status_code, 422)
        self.assertEqual(self.upload("image-only.pdf", b"not a PDF").status_code, 422)

    def test_text_is_bounded_and_names_are_sanitized(self):
        response = self.upload("../../notes.md", b"a" * 10000)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json["name"], "notes.md")
        self.assertEqual(len(response.json["text"]), file_import.MAX_TEXT_CHARS)
        self.assertTrue(response.json["truncated"])

    def test_office_formats_extract_readable_text(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph("Important Word document insight")
        buf = BytesIO()
        doc.save(buf)
        self.assertIn("document insight", self.upload("report.docx", buf.getvalue()).json["text"])

        from openpyxl import Workbook
        book = Workbook()
        book.active.append(["Budget", "Value"])
        book.active.append(["Revenue", 800])
        buf = BytesIO()
        book.save(buf)
        self.assertIn("Revenue | 800", self.upload("analysis.xlsx", buf.getvalue()).json["text"])

        from pptx import Presentation
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = "Strategy document"
        buf = BytesIO()
        deck.save(buf)
        self.assertIn("Strategy document", self.upload("slides.pptx", buf.getvalue()).json["text"])

    def test_job_api_checks_total_attachment_size(self):
        base = {"topic": "Analyze imported files", "workflow": "research", "mode": "team",
                "role": "Researcher", "providers": ["OpenAI"], "demo": True}
        bad = {**base, "attachments": [{"name": "file.txt", "text": "A"*9000}] * 3}
        result = self.client.post("/api/jobs", json=bad, headers=self.headers)
        self.assertEqual(result.status_code, 422)

    def test_practice_never_uses_or_persists_imported_private_text(self):
        marker = "PRIVATE_DOCUMENT_SENTINEL_123"
        payload = {"topic": "Summarize this", "workflow": "research", "mode": "team",
                   "role": "Researcher", "providers": ["OpenAI"], "demo": True,
                   "attachments": [{"name": "secret.txt", "text": marker}]}
        result = self.client.post("/api/jobs", json=payload, headers=self.headers)
        self.assertEqual(result.status_code, 202)
        job_id = result.json["id"]
        for _ in range(40):
            found = self.client.get("/api/jobs/" + job_id, headers=self.headers)
            if found.json["status"] != "running":
                break
            time.sleep(.04)
        self.assertEqual(found.json["status"], "complete")
        self.assertNotIn(marker, found.get_data(as_text=True))
        self.assertNotIn(marker, self.client.get("/api/jobs", headers=self.headers).get_data(as_text=True))

    def test_live_prompt_receives_documents_without_persisting_raw_text(self):
        marker = "PRIVATE_RESEARCH_MARKER"
        payload = {"topic": "Summarize this", "workflow": "research", "mode": "independent",
                   "role": "Researcher", "providers": ["OpenAI"], "demo": False,
                   "attachments": [{"name": "brief.txt", "text": marker}]}
        with patch.dict(os.environ, {"STUDIO_ALLOW_API_USAGE": "true", "STUDIO_DAILY_CALL_CAP": "10",
                                     "OPENAI_API_KEY": "unit-test-fake-never-sent"}):
            with patch.object(self.app_module.threading, "Thread") as thread:
                resp = self.client.post("/api/jobs", json=payload, headers=self.headers)
                self.assertEqual(resp.status_code, 202)
                args = thread.call_args.kwargs["args"]
                self.assertIn(marker, args[1])
                self.assertIn("untrusted source data", args[1])
        jobs = self.client.get("/api/jobs", headers=self.headers)
        self.assertNotIn(marker, jobs.get_data(as_text=True))


if __name__ == "__main__":
    unittest.main()
