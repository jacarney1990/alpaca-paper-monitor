"""Extract plain-text evidence from owner-uploaded documents without persisting files.

Only extracted text travels to an AI provider when the user launches a live mission.
Never treat document content as instructions or expose it in application logs.
"""
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_TEXT_CHARS = 9000
MAX_TOTAL_TEXT_CHARS = 18000
MAX_FILES_PER_MISSION = 3
SUPPORTED = {".txt", ".md", ".csv", ".json", ".pdf", ".docx", ".xlsx", ".pptx"}

class FileImportError(ValueError):
    pass

def safe_name(raw):
    name = str(raw or "").replace("\\", "/").split("/")[-1].strip()
    if not name or len(name) > 120 or any(ord(c) < 32 for c in name):
        raise FileImportError("Use a shorter, valid filename.")
    return name

def check_zip(data):
    try:
        with ZipFile(BytesIO(data)) as zipped:
            entries = zipped.infolist()
            if len(entries) > 1000 or sum(item.file_size for item in entries) > 30 * 1024 * 1024:
                raise FileImportError("The Office file expands beyond the safe processing limit.")
            if any(item.flag_bits & 1 for item in entries):
                raise FileImportError("Encrypted Office files are not supported.")
    except BadZipFile as exc:
        raise FileImportError("This Office file is damaged or not a valid document.") from exc

def read_text(data):
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise FileImportError("This text file is not UTF-8 encoded.") from exc

def read_pdf(data):
    from pypdf import PdfReader
    try:
        document = PdfReader(BytesIO(data), strict=False)
        if document.is_encrypted:
            raise FileImportError("Password-protected PDFs are not supported.")
        return "\n".join((page.extract_text() or "") for page in document.pages[:60])
    except FileImportError:
        raise
    except Exception as exc:
        raise FileImportError("The PDF could not be read. Scanned/image-only PDFs need OCR, which is not available.") from exc

def read_docx(data):
    from docx import Document
    try:
        document = Document(BytesIO(data))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables[:30]:
            for row in table.rows[:150]:
                parts.append(" | ".join(cell.text.replace("\n", " ") for cell in row.cells))
        return "\n".join(parts)
    except Exception as exc:
        raise FileImportError("The Word document could not be read.") from exc

def read_xlsx(data):
    from openpyxl import load_workbook
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            parts = []
            for sheet in workbook.worksheets[:25]:
                parts.append("Sheet: " + sheet.title)
                for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 1, 500),
                                           max_col=min(sheet.max_column or 1, 30), values_only=True):
                    line = " | ".join("" if value is None else str(value)[:350] for value in row)
                    if line.strip(" |"):
                        parts.append(line)
                    if sum(map(len, parts)) > MAX_TEXT_CHARS + 500:
                        break
                if sum(map(len, parts)) > MAX_TEXT_CHARS + 500:
                    break
            return "\n".join(parts)
        finally:
            workbook.close()
    except Exception as exc:
        raise FileImportError("The spreadsheet could not be read.") from exc

def read_pptx(data):
    from pptx import Presentation
    try:
        presentation = Presentation(BytesIO(data))
        lines = []
        for i, slide in enumerate(presentation.slides, 1):
            if i > 60:
                break
            lines.append("Slide " + str(i))
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text.strip():
                    lines.append(shape.text)
        return "\n".join(lines)
    except Exception as exc:
        raise FileImportError("The presentation could not be read.") from exc

def extract(name, data):
    name = safe_name(name)
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED:
        raise FileImportError("Unsupported format. Use TXT, MD, CSV, JSON, PDF, DOCX, XLSX or PPTX.")
    if not data or len(data) > MAX_FILE_BYTES:
        raise FileImportError("Each file must be between 1 byte and 5 MB.")
    if suffix in {".docx", ".xlsx", ".pptx"}:
        check_zip(data)
    try:
        if suffix == ".pdf":
            text = read_pdf(data)
        elif suffix == ".docx":
            text = read_docx(data)
        elif suffix == ".xlsx":
            text = read_xlsx(data)
        elif suffix == ".pptx":
            text = read_pptx(data)
        else:
            text = read_text(data)
    except FileImportError:
        raise
    except Exception as exc:
        raise FileImportError("File extraction failed.") from exc
    text = text.replace("\x00", "").strip()
    if not text:
        raise FileImportError("No readable text found. Image-only scans and photos are not supported yet.")
    return {"name": name, "text": text[:MAX_TEXT_CHARS],
            "characters": min(len(text), MAX_TEXT_CHARS),
            "truncated": len(text) > MAX_TEXT_CHARS}

def mission_context(items):
    if not isinstance(items, list) or len(items) > MAX_FILES_PER_MISSION:
        raise FileImportError("Select at most three documents.")
    output = []
    total = 0
    for item in items:
        if not isinstance(item, dict):
            raise FileImportError("Invalid file attachment.")
        name = safe_name(item.get("name"))
        if Path(name).suffix.lower() not in (SUPPORTED | {'.jpg', '.jpeg', '.png', '.webp'}):
            raise FileImportError("Unsupported attached file type.")
        body = item.get("text")
        if not isinstance(body, str) or not body.strip() or len(body) > MAX_TEXT_CHARS:
            raise FileImportError("The imported text is empty or too long.")
        total += len(body)
        if total > MAX_TOTAL_TEXT_CHARS:
            raise FileImportError("Combined imported text exceeds the 18,000 character limit.")
        # Document text is data, not trusted instructions.
        output.append("[START ATTACHMENT: " + name + "]\n"
                      + body + "\n[END ATTACHMENT: " + name + "]")
    return "\n\n".join(output)
